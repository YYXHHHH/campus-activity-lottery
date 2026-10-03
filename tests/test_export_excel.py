"""FR-7.3~7.6 Excel 导出。"""

from __future__ import annotations

import io
from urllib.parse import unquote

from openpyxl import load_workbook
from sqlalchemy import select

from app.models import Registration
from helpers import create_activity, create_students, make_user, run_lottery


def _load(response):
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    return load_workbook(io.BytesIO(response.content))


def _filename(response):
    disposition = response.headers["content-disposition"]
    marker = "filename*=UTF-8''"
    assert marker in disposition, disposition
    return unquote(disposition.split(marker, 1)[1])


def _prepare(client, db, quota=2, count=5, title="摄影采风活动"):
    org = make_user(client, "org1", "ORGANIZER")
    activity = create_activity(client, org["headers"], quota=quota, title=title)
    students = create_students(client, count)
    for index, student in enumerate(students):
        client.post(
            f"/api/activities/{activity['id']}/registrations",
            json={"accept_waitlist": index < 4},
            headers=student["headers"],
        )
    run_lottery(client, org["headers"], activity["id"])
    return org, activity, students


def _winners(client, org, activity_id):
    return client.get(f"/api/activities/{activity_id}/winners", headers=org["headers"]).json()["data"]["winners"]


def test_export_registration_sheet(client, db):
    org, activity, _students = _prepare(client, db)
    workbook = _load(client.get(f"/api/activities/{activity['id']}/export", params={"type": "all"}, headers=org["headers"]))

    assert _filename(client.get(f"/api/activities/{activity['id']}/export", params={"type": "all"}, headers=org["headers"])).startswith(
        "摄影采风活动_报名表_"
    )

    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    assert headers == ["序号", "学号/工号", "姓名", "报名时间", "接受候补", "状态", "抽签序号", "签到状态", "签到时间"]
    assert sheet.freeze_panes == "A2"
    assert all(cell.font.bold for cell in sheet[1])
    rows = [row for row in sheet.iter_rows(min_row=2, values_only=True) if row[0] is not None]
    assert len(rows) == 5
    assert {row[5] for row in rows} <= {"已中签", "候补中", "未中签", "待抽签", "已退出"}
    assert rows[0][1].startswith("stu")
    assert rows[0][3].endswith("Z")  # 报名时间为 ISO UTC
    assert rows[0][4] in ("是", "否")


def test_export_winners_sheet_contains_qr_code(client, db):
    org, activity, students = _prepare(client, db)
    response = client.get(f"/api/activities/{activity['id']}/export", params={"type": "winners"}, headers=org["headers"])
    workbook = _load(response)
    assert _filename(response).startswith("摄影采风活动_中签表_")

    headers = [cell.value for cell in workbook.active[1]]
    assert headers == ["序号", "学号/工号", "姓名", "抽签序号", "签到码", "签到状态", "签到时间"]
    rows = [row for row in workbook.active.iter_rows(min_row=2, values_only=True) if row[0] is not None]
    assert len(rows) == 2
    assert all(row[4] for row in rows)  # 中签者均有签到码
    assert workbook.active.freeze_panes == "A2"

    # 已退出者不出现在中签表
    winner_names = {row[1] for row in rows}
    leaver_username = next(name for name in winner_names)
    leaver = next(s for s in students if s["username"] == leaver_username)
    client.delete(f"/api/activities/{activity['id']}/registrations/me", headers=leaver["headers"])

    after_rows = [
        row
        for row in _load(
            client.get(f"/api/activities/{activity['id']}/export", params={"type": "winners"}, headers=org["headers"])
        ).active.iter_rows(min_row=2, values_only=True)
        if row[0] is not None
    ]
    assert leaver_username not in {row[1] for row in after_rows}


def test_export_empty_activity(client):
    org = make_user(client, "org1", "ORGANIZER")
    activity = create_activity(client, org["headers"])
    workbook = _load(client.get(f"/api/activities/{activity['id']}/export", params={"type": "all"}, headers=org["headers"]))
    rows = [row for row in workbook.active.iter_rows(min_row=2, values_only=True) if row[0] is not None]
    assert rows == []
    assert workbook.active["A1"].value == "序号"


def test_export_type_validation(client):
    org = make_user(client, "org1", "ORGANIZER")
    activity = create_activity(client, org["headers"])
    bad = client.get(f"/api/activities/{activity['id']}/export", params={"type": "pdf"}, headers=org["headers"])
    assert bad.status_code == 400 and bad.json()["code"] == 40001
