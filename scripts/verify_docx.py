# -*- coding: utf-8 -*-
"""Verify the generated Word documents under docs/: structure stats + key-content assertions."""
import os, glob
from docx import Document

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOC = os.path.join(BASE, 'docs')
files = [f for f in sorted(glob.glob(os.path.join(DOC, '**', '*.docx'), recursive=True))
         if '99-archive' not in f]
print('docx files:', len(files))
for f in files:
    d = Document(f)
    paras = [p.text for p in d.paragraphs if p.text.strip()]
    rows = sum(len(t.rows) for t in d.tables)
    title = paras[0][:38] if paras else '(EMPTY)'
    print('%-56s paras=%4d tables=%3d rows=%4d %6.1fKB | %s' % (
        os.path.relpath(f, DOC), len(paras), len(d.tables), rows, os.path.getsize(f) / 1024, title))

checks = {
    '01-project-plan/project-development-plan.docx': ['里程碑', '风险'],
    '01-project-plan/software-quality-assurance-plan.docx': ['质量目标', '缺陷管理'],
    '01-project-plan/software-configuration-management-plan.docx': ['配置项识别', '基线'],
    '02-requirements/requirements-specification.docx': ['功能需求', '权限矩阵', 'FR-7'],
    '03-design/high-level-design.docx': ['分层架构', 'M10 前端页面'],
    '03-design/detailed-design.docx': ['抽签算法', '模块间依赖'],
    '03-design/database-design.docx': ['表结构 DDL', '数据完整性'],
    '03-design/interface-design.docx': ['全局约定', '签到'],
    '04-implementation/implementation-notes.docx': ['交付物总览', '集成点'],
    '05-testing/test-plan.docx': ['自动化测试', '附录 D'],
    '05-testing/test-report.docx': ['执行环境', '缺陷'],
    '06-deployment/deployment-manual.docx': ['部署', '排障'],
    '06-deployment/user-manual.docx': ['学生', '组织者'],
    '07-acceptance/acceptance-report.docx': ['A1', '验收'],
    '07-acceptance/project-summary.docx': ['目标达成', '改进'],
    'documentation-index.docx': ['Standard document map', 'Requirements'],
}
print('\n-- content assertions --')
ok = True
for rel, needles in checks.items():
    p = os.path.join(DOC, rel)
    if not os.path.exists(p):
        print('MISSING', rel); ok = False; continue
    d = Document(p)
    alltext = '\n'.join(x.text for x in d.paragraphs)
    for tb in d.tables:
        for r in tb.rows:
            for cc in r.cells:
                alltext += '\n' + cc.text
    miss = [n for n in needles if n not in alltext]
    print(('OK  ' if not miss else 'FAIL') + ' ' + rel + ('' if not miss else ' missing=' + str(miss)))
    ok = ok and not miss
print('\nRESULT', 'PASS' if ok else 'FAIL')
