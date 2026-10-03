"""M7 Excel 导出（§5.8、S9）。"""

from __future__ import annotations

import io
import logging

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import joinedload

from app.database import iso_z, utcnow
from app.models import (
    REGISTRATION_STATUS_CN,
    Activity,
    Checkin,
    Registration,
    RegistrationStatus,
    User,
)

logger = logging.getLogger(__name__)

REGISTRATION_HEADERS = ["序号", "学号/工号", "姓名", "报名时间", "接受候补", "状态", "抽签序号", "签到状态", "签到时间"]
WINNER_HEADERS = ["序号", "学号/工号", "姓名", "抽签序号", "签到码", "签到状态", "签到时间"]


def _fetch_rows(db, activity: Activity, winners_only: bool):
    filters = [Registration.activity_id == activity.id]
    if winners_only:
        filters.append(Registration.status == RegistrationStatus.WON)
    else:
        filters.append(Registration.status != RegistrationStatus.CANCELLED)
    return db.scalars(
        select(Registration)
        .options(joinedload(Registration.user), joinedload(Registration.checkin))
        .where(*filters)
        .order_by(Registration.lottery_rank.asc() if winners_only else Registration.id.asc())
    ).unique().all()


def _cell_width(value: str) -> int:
    return sum(2 if ord(char) > 0x2E80 else 1 for char in str(value))


def _build_sheet(ws, headers: list[str], rows: list[list]) -> None:
    ws.append(headers)
    for header_cell in ws[1]:
        header_cell.font = Font(bold=True)
        header_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"

    for row in rows:
        ws.append(row)

    widths = [_cell_width(header) for header in headers]
    for row in rows:
        for index, value in enumerate(row):
            widths[index] = max(widths[index], _cell_width(value if value is not None else ""))
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = min(max(width + 4, 10), 60)


def build_workbook(db, activity: Activity, export_type: str) -> tuple[bytes, str]:
    """返回 (xlsx 字节流, 文件名)。"""
    winners_only = export_type == "winners"
    registrations = _fetch_rows(db, activity, winners_only)

    rows: list[list] = []
    for index, registration in enumerate(registrations, start=1):
        user: User = registration.user
        checkin: Checkin | None = registration.checkin
        status_cn = REGISTRATION_STATUS_CN.get(str(registration.status), str(registration.status))
        checked_in_cn = "已签到" if checkin else "未签到"
        checked_at = iso_z(checkin.checked_in_at) if checkin else ""
        if winners_only:
            rows.append(
                [index, user.username, user.real_name, registration.lottery_rank, registration.qr_token or "", checked_in_cn, checked_at]
            )
        else:
            rows.append(
                [
                    index,
                    user.username,
                    user.real_name,
                    iso_z(registration.created_at),
                    "是" if registration.accept_waitlist else "否",
                    status_cn,
                    registration.lottery_rank if registration.lottery_rank is not None else "",
                    checked_in_cn,
                    checked_at,
                ]
            )

    wb = Workbook()
    ws = wb.active
    ws.title = "中签名单" if winners_only else "报名名单"
    _build_sheet(ws, WINNER_HEADERS if winners_only else REGISTRATION_HEADERS, rows)

    buffer = io.BytesIO()
    wb.save(buffer)
    suffix = "中签表" if winners_only else "报名表"
    filename = f"{activity.title}_{suffix}_{utcnow().astimezone():%Y%m%d}.xlsx"
    logger.info(
        "[export] 导出 activity_id=%s type=%s row_count=%s operator_id=%s",
        activity.id,
        export_type,
        len(rows),
        activity.organizer_id,
    )
    return buffer.getvalue(), filename


def content_disposition(filename: str) -> str:
    """RFC 5987 编码中文文件名（S9）。"""
    from urllib.parse import quote

    fallback = "export.xlsx"
    return f"attachment; filename={fallback}; filename*=UTF-8''{quote(filename)}"
